"""Document loaders.

Each loader returns one or more LangChain ``Document`` objects carrying rich
metadata (filename, source path, document type, title, page number where
available, and for DOCX an inferred section).
"""

from __future__ import annotations

import logging
import re
from pathlib import Path

from langchain_core.documents import Document

from app.exceptions import UnsupportedFileTypeError

logger = logging.getLogger(__name__)

VALID_EXTENSIONS = {".pdf", ".docx", ".txt"}

_HEADING_RE = re.compile(r"^(?:\d+(?:\.\d+)*[\.\)]\s*)?[A-Z][A-Za-z0-9&' ,/\-]{3,}$")


def _title_from_text(text: str) -> str:
    """Best-effort extraction of a human readable section title."""
    for line in text.splitlines():
        line = line.strip()
        if line and _HEADING_RE.match(line) and len(line) <= 120:
            return line
    return "General"


def _detect_section(text: str) -> str:
    """Infer a section name from the first heading-like line in a chunk."""
    for line in text.splitlines():
        line = line.strip()
        if line and _HEADING_RE.match(line) and len(line) <= 120:
            return line
    return "General"


def _base_metadata(path: Path, doc_type: str, title: str | None = None) -> dict:
    return {
        "filename": path.name,
        "source": str(path.resolve()),
        "doc_type": doc_type,
        "title": title or _title_from_text(path.stem.replace("_", " ")),
    }


def _clean_text(text: str) -> str:
    """Normalise whitespace without destroying paragraph boundaries."""
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)  # collapse runs of spaces/tabs
    text = re.sub(r"\n{3,}", "\n\n", text)  # collapse excessive blank lines
    return text.strip()


# --------------------------------------------------------------------------- #
# Individual loaders
# --------------------------------------------------------------------------- #
def _load_pdf(path: Path) -> list[Document]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    docs: list[Document] = []
    for index, page in enumerate(reader.pages, start=1):
        raw = page.extract_text() or ""
        text = _clean_text(raw)
        if not text:
            logger.debug("Skipped empty page %d of %s", index, path.name)
            continue
        metadata = _base_metadata(path, "pdf", title=_title_from_text(text))
        metadata["page"] = index
        metadata["num_pages"] = len(reader.pages)
        docs.append(Document(page_content=text, metadata=metadata))
    if not docs:
        logger.warning("No extractable text found in PDF '%s'.", path.name)
    return docs


def _load_docx(path: Path) -> list[Document]:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    lines: list[str] = []
    section = "General"
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower()
        if style.startswith("heading"):
            if lines:
                lines.append("\n")
            lines.append(text)
            section = text
        else:
            lines.append(text)

    # Include table content so tabular policies (e.g. leave tables) are indexed.
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))

    text = _clean_text("\n".join(lines))
    metadata = _base_metadata(path, "docx", title=_title_from_text(text))
    metadata["section"] = section if section != "General" else None
    return [Document(page_content=text, metadata=metadata)] if text else []


def _load_txt(path: Path) -> list[Document]:
    raw = path.read_bytes()
    for encoding in ("utf-8", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - extremely unlikely after latin-1
        text = raw.decode("utf-8", errors="replace")

    text = _clean_text(text)
    if not text:
        logger.warning("TXT file '%s' is empty.", path.name)
        return []
    metadata = _base_metadata(path, "txt", title=_title_from_text(text))
    return [Document(page_content=text, metadata=metadata)]


# --------------------------------------------------------------------------- #
# Public API
# --------------------------------------------------------------------------- #
def load_document(path: Path) -> list[Document]:
    """Load a single supported document into LangChain ``Document`` objects."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _load_pdf(path)
    if ext == ".docx":
        return _load_docx(path)
    if ext == ".txt":
        return _load_txt(path)
    raise UnsupportedFileTypeError(
        f"Unsupported file type '{ext}' for '{path.name}'. "
        f"Supported: {', '.join(sorted(VALID_EXTENSIONS))}."
    )


def normalize_metadata(metadata: dict) -> dict:
    """Ensure only Chroma-safe scalar values remain in metadata."""
    return {
        k: v
        for k, v in metadata.items()
        if v is not None and isinstance(v, (str, int, float, bool))
    }


__all__ = [
    "load_document",
    "normalize_metadata",
    "_detect_section",
    "VALID_EXTENSIONS",
]